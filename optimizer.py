"""
Steel S&OP / MPS optimization agent.

This script uses only the data supplied in the Excel template:
- Sinter Plant annual capacity
- Blast Furnace annual capacity
- Rolling Mill 1 and Rolling Mill 2 SKU capacities
- Rolling-mill setup/changeover times

Demand is intentionally user-editable. On the first run, the script creates an
input workbook with a temporary six-month demand scenario and an Active Y/N
toggle for every SKU-month cell. Edit those cells, rerun the script, and the
optimized MPS/output workbook is regenerated.

Raw-material BOM factors are not supplied in the template, so the script does
not invent them. If you fill the optional BOM input sheet, a raw-material
requirements table is produced; otherwise the MRP-style output is limited to
finished-goods production and capacity requirements.
"""

from __future__ import annotations

import argparse
import itertools
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


MONTHS = [f"Month {i}" for i in range(1, 7)]
TEMP_DEMAND = {
    "SKU 1": [70, 75, 80, 80, 85, 90],
    "SKU 2": [70, 75, 75, 80, 85, 85],
    "SKU 3": [90, 95, 100, 100, 105, 110],
    "SKU 4": [60, 65, 70, 65, 70, 75],
    "SKU 5": [90, 85, 80, 75, 65, 50],
}


@dataclass
class SourceData:
    skus: List[str]
    months: List[str]
    sinter_annual: float
    blast_furnace_annual: float
    mill_caps_annual: Dict[str, Dict[str, float]]
    setup_adjacent_hours: float
    setup_non_adjacent_hours: float


@dataclass
class MonthPlan:
    month: str
    sku: str
    demand: float
    active: bool
    effective_demand: float
    rm1_qty: float
    rm2_qty: float
    planned_production: float
    unmet_demand: float


@dataclass
class CampaignPlan:
    month: str
    mill: str
    active_skus: Tuple[str, ...]
    sequence: Tuple[str, ...]
    setup_hours: float
    normalized_load: float


def to_float(value, default: Optional[float] = None) -> Optional[float]:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"[-+]?\d*\.?\d+", str(value))
    return float(match.group()) if match else default


def norm_label(value) -> str:
    return str(value or "").strip().lower()


def parse_template(path: Path) -> SourceData:
    wb = load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))

    def find_row_containing(text: str) -> int:
        wanted = text.strip().lower()
        for idx, row in enumerate(rows, start=1):
            if any(norm_label(cell) == wanted for cell in row):
                return idx
        raise ValueError(f"Could not find '{text}' in the template.")

    demand_title_row = find_row_containing("Demand")
    header_row = demand_title_row + 1
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    product_col = None
    for c, value in enumerate(headers, start=1):
        if norm_label(value) == "product":
            product_col = c
            break
    if product_col is None:
        raise ValueError("Could not find the Product header in the demand section.")

    months = []
    month_cols = []
    for c in range(product_col + 1, ws.max_column + 1):
        value = str(ws.cell(header_row, c).value or "").strip()
        if value.lower().startswith("month"):
            months.append(value)
            month_cols.append(c)
    if not months:
        months = MONTHS

    skus = []
    r = header_row + 1
    while r <= ws.max_row:
        value = ws.cell(r, product_col).value
        if value is None or str(value).strip() == "":
            break
        skus.append(str(value).strip())
        r += 1
    if not skus:
        raise ValueError("No SKU rows were found in the demand section.")

    def value_right_of(label: str) -> float:
        row_num = find_row_containing(label)
        for c in range(1, ws.max_column + 1):
            if norm_label(ws.cell(row_num, c).value) == label.lower():
                for j in range(c + 1, ws.max_column + 1):
                    value = to_float(ws.cell(row_num, j).value)
                    if value is not None:
                        return value
        raise ValueError(f"Could not find numeric value for '{label}'.")

    sinter = value_right_of("Sinter plant")
    blast = value_right_of("Blast Furnace")

    rm_rows = {}
    for r in range(1, ws.max_row + 1):
        row_text = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)).lower()
        if "rolling mill 1" in row_text:
            rm_rows["Rolling Mill 1"] = r
        if "rolling mill 2" in row_text:
            rm_rows["Rolling Mill 2"] = r
    if len(rm_rows) != 2:
        raise ValueError("Could not find both rolling mill rows in the template.")

    sku_header_row = min(rm_rows.values()) - 1
    sku_cols = {}
    for c in range(1, ws.max_column + 1):
        label = str(ws.cell(sku_header_row, c).value or "").replace(" ", "").upper()
        if label.startswith("SKU"):
            sku_cols[label] = c

    mill_caps = {"Rolling Mill 1": {}, "Rolling Mill 2": {}}
    for mill, row_num in rm_rows.items():
        for sku in skus:
            key = sku.replace(" ", "").upper()
            if key not in sku_cols:
                raise ValueError(f"Could not find capacity column for {sku}.")
            value = to_float(ws.cell(row_num, sku_cols[key]).value)
            if value is None:
                raise ValueError(f"Missing capacity for {mill} / {sku}.")
            mill_caps[mill][sku] = value

    adjacent = value_right_of("SKU (number) to (number +/- 1)")
    non_adjacent = value_right_of("SKU (number) to (>number +/- 1)")

    return SourceData(
        skus=skus,
        months=months,
        sinter_annual=sinter,
        blast_furnace_annual=blast,
        mill_caps_annual=mill_caps,
        setup_adjacent_hours=adjacent,
        setup_non_adjacent_hours=non_adjacent,
    )


def style_sheet(ws):
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color="D9E2EC")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            if cell.value is not None:
                cell.border = Border(bottom=thin)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    calc_fill = PatternFill("solid", fgColor="E2F0D9")

    for row in ws.iter_rows():
        for cell in row:
            if cell.row == 1:
                cell.font = Font(bold=True, size=14, color="1F2937")
            elif isinstance(cell.value, str) and cell.value in {
                "SKU",
                "Month",
                "Mill",
                "Demand",
                "Active?",
                "Effective Demand",
                "Planned Production",
            }:
                cell.fill = header_fill
                cell.font = header_font
            elif cell.row in {3, 12, 22, 32}:
                cell.fill = sub_fill
                cell.font = Font(bold=True)

    for ws_col in range(1, ws.max_column + 1):
        col_letter = get_column_letter(ws_col)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 11), 24)

    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 21

    return input_fill, calc_fill


def create_input_workbook(source: SourceData, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Demand_Input"
    ws["A1"] = "Steel S&OP Demand Input"
    ws["A3"] = "Edit yellow demand cells and Y/N active toggles, then rerun the Python agent."

    ws.cell(5, 1, "Demand")
    for c, month in enumerate(source.months, start=2):
        ws.cell(5, c, month)
    for r, sku in enumerate(source.skus, start=6):
        ws.cell(r, 1, sku)
        values = TEMP_DEMAND.get(sku, [0] * len(source.months))
        for c, value in enumerate(values, start=2):
            ws.cell(r, c, value)

    toggle_start = 14
    ws.cell(toggle_start, 1, "Active?")
    for c, month in enumerate(source.months, start=2):
        ws.cell(toggle_start, c, month)
    for r, sku in enumerate(source.skus, start=toggle_start + 1):
        ws.cell(r, 1, sku)
        for c in range(2, 2 + len(source.months)):
            ws.cell(r, c, "Y")

    ws.cell(23, 1, "Given Capacities")
    ws.cell(24, 1, "Stage")
    ws.cell(24, 2, "Annual Capacity")
    ws.cell(24, 3, "Monthly Capacity")
    ws.cell(25, 1, "Sinter Plant")
    ws.cell(25, 2, source.sinter_annual)
    ws.cell(25, 3, source.sinter_annual / 12)
    ws.cell(26, 1, "Blast Furnace")
    ws.cell(26, 2, source.blast_furnace_annual)
    ws.cell(26, 3, source.blast_furnace_annual / 12)

    ws.cell(29, 1, "Rolling Mill Capacities")
    ws.cell(30, 1, "Mill")
    for c, sku in enumerate(source.skus, start=2):
        ws.cell(30, c, sku)
    for r, mill in enumerate(["Rolling Mill 1", "Rolling Mill 2"], start=31):
        ws.cell(r, 1, mill)
        for c, sku in enumerate(source.skus, start=2):
            ws.cell(r, c, source.mill_caps_annual[mill][sku])

    ws.cell(35, 1, "Setup Times")
    ws.cell(36, 1, "Adjacent SKU change")
    ws.cell(36, 2, source.setup_adjacent_hours)
    ws.cell(37, 1, "Non-adjacent SKU change")
    ws.cell(37, 2, source.setup_non_adjacent_hours)

    bom = wb.create_sheet("BOM_Input_Optional")
    bom["A1"] = "Optional Raw-Material BOM Inputs"
    bom["A3"] = "Material"
    bom["B3"] = "Consumption per ton finished steel"
    bom["C3"] = "Unit"
    bom["A4"] = "Leave blank unless your assignment supplies BOM factors."
    bom["B4"] = None
    bom["C4"] = "tons / ton"

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Model Boundary"
    notes["A3"] = "Temporary demand is provided only to make the optimizer runnable."
    notes["A4"] = "True raw-material MRP requires BOM/consumption factors, which are not supplied in the template."
    notes["A5"] = "Rolling-mill capacity is modeled using normalized load: sum(production / monthly SKU capacity) <= 1."
    notes["A6"] = "Setup times are minimized and reported as campaign hours; they are not converted into lost tonnage because operating hours are not supplied."

    input_fill, _ = style_sheet(ws)
    for row in range(6, 6 + len(source.skus)):
        for col in range(2, 2 + len(source.months)):
            ws.cell(row, col).fill = input_fill
    for row in range(toggle_start + 1, toggle_start + 1 + len(source.skus)):
        for col in range(2, 2 + len(source.months)):
            ws.cell(row, col).fill = input_fill

    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"B{toggle_start + 1}:{get_column_letter(1 + len(source.months))}{toggle_start + len(source.skus)}")
    ws.freeze_panes = "B6"
    style_sheet(bom)
    style_sheet(notes)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_inputs(input_path: Path, source: SourceData) -> Tuple[Dict[str, Dict[str, float]], Dict[str, Dict[str, bool]], Dict[str, float]]:
    wb = load_workbook(input_path, data_only=True)
    ws = wb["Demand_Input"]

    demand = {sku: {} for sku in source.skus}
    active = {sku: {} for sku in source.skus}

    for r, sku in enumerate(source.skus, start=6):
        for c, month in enumerate(source.months, start=2):
            demand[sku][month] = max(0.0, to_float(ws.cell(r, c).value, 0.0) or 0.0)

    toggle_start = 15
    for r, sku in enumerate(source.skus, start=toggle_start):
        for c, month in enumerate(source.months, start=2):
            active[sku][month] = str(ws.cell(r, c).value or "Y").strip().upper() != "N"

    bom_factors = {}
    if "BOM_Input_Optional" in wb.sheetnames:
        bom = wb["BOM_Input_Optional"]
        for r in range(4, bom.max_row + 1):
            material = bom.cell(r, 1).value
            factor = to_float(bom.cell(r, 2).value)
            if material and factor is not None and factor > 0:
                bom_factors[str(material).strip()] = factor

    return demand, active, bom_factors


def sku_number(sku: str) -> int:
    match = re.search(r"\d+", sku)
    if not match:
        raise ValueError(f"SKU label must contain a number: {sku}")
    return int(match.group())


def transition_hours(a: str, b: str, adjacent: float, non_adjacent: float) -> float:
    return adjacent if abs(sku_number(a) - sku_number(b)) == 1 else non_adjacent


def best_sequence(skus: Sequence[str], adjacent: float, non_adjacent: float) -> Tuple[Tuple[str, ...], float]:
    if len(skus) <= 1:
        return tuple(skus), 0.0
    best = None
    for perm in itertools.permutations(skus):
        hours = sum(transition_hours(perm[i], perm[i + 1], adjacent, non_adjacent) for i in range(len(perm) - 1))
        if best is None or hours < best[1]:
            best = (perm, hours)
    return best


def all_subsets(items: Sequence[str]) -> Iterable[Tuple[str, ...]]:
    for size in range(0, len(items) + 1):
        for subset in itertools.combinations(items, size):
            yield subset


def allocate_for_sets(
    demand: Dict[str, float],
    set1: Tuple[str, ...],
    set2: Tuple[str, ...],
    cap1: Dict[str, float],
    cap2: Dict[str, float],
) -> Optional[Tuple[Dict[str, float], Dict[str, float], float, float]]:
    x1 = {}
    lower = {}
    upper = {}
    active_demand_skus = [sku for sku, qty in demand.items() if qty > 1e-9]
    for sku in active_demand_skus:
        if sku not in set1 and sku not in set2:
            return None
        lo = demand[sku] if sku not in set2 else 0.0
        hi = 0.0 if sku not in set1 else demand[sku]
        if lo > hi + 1e-9:
            return None
        lower[sku] = lo
        upper[sku] = hi
        x1[sku] = lo

    def load(mill: int) -> float:
        if mill == 1:
            return sum(x1.get(sku, 0.0) / cap1[sku] for sku in active_demand_skus if cap1[sku] > 0)
        return sum((demand[sku] - x1.get(sku, 0.0)) / cap2[sku] for sku in active_demand_skus if cap2[sku] > 0)

    if load(1) > 1.0 + 1e-9:
        return None

    if load(2) > 1.0 + 1e-9:
        needed = load(2) - 1.0
        candidates = []
        for sku in active_demand_skus:
            room = upper[sku] - x1[sku]
            if room > 1e-9:
                benefit = 1.0 / cap2[sku]
                cost = 1.0 / cap1[sku]
                candidates.append((benefit / cost, sku, room, benefit, cost))
        candidates.sort(reverse=True)
        for _, sku, room, benefit, _ in candidates:
            if needed <= 1e-9:
                break
            move = min(room, needed / benefit)
            x1[sku] += move
            needed -= move * benefit
        if needed > 1e-7:
            return None

    l1 = load(1)
    l2 = load(2)
    if l1 > 1.0 + 1e-7 or l2 > 1.0 + 1e-7:
        return None

    rm1 = {sku: round(x1.get(sku, 0.0), 6) for sku in demand}
    rm2 = {sku: round(demand[sku] - rm1[sku], 6) for sku in demand}
    return rm1, rm2, l1, l2


def optimize_month(source: SourceData, month_demand: Dict[str, float]) -> Tuple[Dict[str, float], Dict[str, float], CampaignPlan, CampaignPlan, float]:
    monthly_caps = {
        mill: {sku: source.mill_caps_annual[mill][sku] / 12.0 for sku in source.skus}
        for mill in ["Rolling Mill 1", "Rolling Mill 2"]
    }
    active_skus = tuple(sku for sku in source.skus if month_demand[sku] > 1e-9)
    if not active_skus:
        empty1 = CampaignPlan("", "Rolling Mill 1", tuple(), tuple(), 0.0, 0.0)
        empty2 = CampaignPlan("", "Rolling Mill 2", tuple(), tuple(), 0.0, 0.0)
        return {sku: 0.0 for sku in source.skus}, {sku: 0.0 for sku in source.skus}, empty1, empty2, 0.0

    best = None
    for set1 in all_subsets(active_skus):
        for set2 in all_subsets(active_skus):
            allocation = allocate_for_sets(month_demand, set1, set2, monthly_caps["Rolling Mill 1"], monthly_caps["Rolling Mill 2"])
            if allocation is None:
                continue
            rm1, rm2, load1, load2 = allocation
            seq1, setup1 = best_sequence(set1, source.setup_adjacent_hours, source.setup_non_adjacent_hours)
            seq2, setup2 = best_sequence(set2, source.setup_adjacent_hours, source.setup_non_adjacent_hours)
            setup_total = setup1 + setup2
            balance_penalty = abs(load1 - load2)
            load_total = load1 + load2
            score = (setup_total, balance_penalty, load_total)
            if best is None or score < best[0]:
                best = (score, rm1, rm2, seq1, setup1, load1, seq2, setup2, load2)

    if best is None:
        raise ValueError(
            "Demand cannot be met with the supplied rolling-mill capacities for at least one month. "
            "Reduce demand or toggle off some SKU-month cells."
        )

    _, rm1, rm2, seq1, setup1, load1, seq2, setup2, load2 = best
    camp1 = CampaignPlan("", "Rolling Mill 1", tuple(s for s in active_skus if rm1.get(s, 0) > 1e-9), seq1, setup1, load1)
    camp2 = CampaignPlan("", "Rolling Mill 2", tuple(s for s in active_skus if rm2.get(s, 0) > 1e-9), seq2, setup2, load2)
    return rm1, rm2, camp1, camp2, setup1 + setup2


def run_optimization(source: SourceData, demand: Dict[str, Dict[str, float]], active: Dict[str, Dict[str, bool]]):
    plans: List[MonthPlan] = []
    campaigns: List[CampaignPlan] = []
    capacity_rows = []

    upstream_monthly = {
        "Sinter Plant": source.sinter_annual / 12.0,
        "Blast Furnace": source.blast_furnace_annual / 12.0,
    }

    for month in source.months:
        month_demand = {
            sku: demand[sku][month] if active[sku][month] else 0.0
            for sku in source.skus
        }
        total_required = sum(month_demand.values())
        upstream_cap = min(upstream_monthly.values())
        if total_required > upstream_cap + 1e-7:
            raise ValueError(
                f"{month} demand is {total_required:.2f} tons, above the supplied monthly bottleneck "
                f"capacity of {upstream_cap:.2f} tons. Reduce demand or toggle cells off."
            )

        rm1, rm2, camp1, camp2, setup_total = optimize_month(source, month_demand)
        camp1.month = month
        camp2.month = month
        campaigns.extend([camp1, camp2])

        for sku in source.skus:
            planned = rm1[sku] + rm2[sku]
            plans.append(
                MonthPlan(
                    month=month,
                    sku=sku,
                    demand=demand[sku][month],
                    active=active[sku][month],
                    effective_demand=month_demand[sku],
                    rm1_qty=rm1[sku],
                    rm2_qty=rm2[sku],
                    planned_production=planned,
                    unmet_demand=max(0.0, month_demand[sku] - planned),
                )
            )

        capacity_rows.append(
            {
                "Month": month,
                "Effective Demand": total_required,
                "Sinter Monthly Capacity": upstream_monthly["Sinter Plant"],
                "Blast Furnace Monthly Capacity": upstream_monthly["Blast Furnace"],
                "Sinter Utilization": total_required / upstream_monthly["Sinter Plant"] if upstream_monthly["Sinter Plant"] else 0,
                "Blast Furnace Utilization": total_required / upstream_monthly["Blast Furnace"] if upstream_monthly["Blast Furnace"] else 0,
                "RM1 Normalized Load": camp1.normalized_load,
                "RM2 Normalized Load": camp2.normalized_load,
                "Setup Hours": setup_total,
            }
        )
    return plans, campaigns, capacity_rows


def add_rows(ws, headers: Sequence[str], rows: Sequence[Sequence], start_row: int = 3):
    for c, h in enumerate(headers, start=1):
        ws.cell(start_row, c, h)
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)


def write_output(
    output_path: Path,
    source: SourceData,
    input_path: Path,
    demand: Dict[str, Dict[str, float]],
    active: Dict[str, Dict[str, bool]],
    bom_factors: Dict[str, float],
    plans: List[MonthPlan],
    campaigns: List[CampaignPlan],
    capacity_rows: List[Dict[str, float]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = "Steel S&OP Optimization Dashboard"
    ws["A3"] = "Input workbook"
    ws["B3"] = str(input_path.name)
    total_demand = sum(p.effective_demand for p in plans)
    total_production = sum(p.planned_production for p in plans)
    total_unmet = sum(p.unmet_demand for p in plans)
    total_setup = sum(c.setup_hours for c in campaigns)
    bottleneck_cap = source.blast_furnace_annual / 2.0
    metrics = [
        ("Six-month effective demand", total_demand, "tons"),
        ("Optimized planned production", total_production, "tons"),
        ("Unmet demand", total_unmet, "tons"),
        ("Six-month BF capacity", bottleneck_cap, "tons"),
        ("BF utilization", total_demand / bottleneck_cap if bottleneck_cap else 0, "%"),
        ("Reported setup/changeover time", total_setup, "hours"),
    ]
    add_rows(ws, ["Metric", "Value", "Unit"], metrics, 5)

    ws2 = wb.create_sheet("Demand_Used")
    headers = ["SKU"] + source.months + ["Six-month Total"]
    rows = []
    for sku in source.skus:
        row = [sku] + [demand[sku][m] if active[sku][m] else 0.0 for m in source.months]
        row.append(sum(row[1:]))
        rows.append(row)
    add_rows(ws2, headers, rows)

    ws3 = wb.create_sheet("Optimized_MPS")
    add_rows(
        ws3,
        ["Month", "SKU", "Input Demand", "Active?", "Effective Demand", "RM1 Qty", "RM2 Qty", "Planned Production", "Unmet Demand"],
        [
            [p.month, p.sku, p.demand, "Y" if p.active else "N", p.effective_demand, p.rm1_qty, p.rm2_qty, p.planned_production, p.unmet_demand]
            for p in plans
        ],
    )

    ws4 = wb.create_sheet("Campaign_Plan")
    add_rows(
        ws4,
        ["Month", "Mill", "Active SKUs", "Recommended Sequence", "Setup Hours", "Normalized Capacity Load"],
        [
            [c.month, c.mill, ", ".join(c.active_skus) or "No production", " -> ".join(c.sequence) or "No production", c.setup_hours, c.normalized_load]
            for c in campaigns
        ],
    )

    ws5 = wb.create_sheet("Capacity_Requirements")
    add_rows(
        ws5,
        list(capacity_rows[0].keys()),
        [[row[key] for key in capacity_rows[0].keys()] for row in capacity_rows],
    )

    ws6 = wb.create_sheet("MRP_Production_Req")
    add_rows(
        ws6,
        ["Month", "SKU", "Gross Requirement", "Planned Production Receipt", "Ending Finished Inventory"],
        [[p.month, p.sku, p.effective_demand, p.planned_production, max(0.0, p.planned_production - p.effective_demand)] for p in plans],
    )

    ws7 = wb.create_sheet("Raw_Material_MRP")
    if bom_factors:
        rows = []
        monthly_prod = {m: sum(p.planned_production for p in plans if p.month == m) for m in source.months}
        for material, factor in bom_factors.items():
            for month in source.months:
                rows.append([material, month, factor, monthly_prod[month], factor * monthly_prod[month]])
        add_rows(ws7, ["Material", "Month", "Consumption per Ton", "Finished Production", "Gross Material Requirement"], rows)
    else:
        add_rows(
            ws7,
            ["Status", "Reason", "How to enable"],
            [[
                "Not computed",
                "The supplied template does not include raw-material BOM/consumption factors.",
                "Fill BOM_Input_Optional in the input workbook, then rerun this script.",
            ]],
        )

    ws8 = wb.create_sheet("Given_Source_Data")
    rows = [
        ["Sinter Plant annual capacity", source.sinter_annual, "tons/year"],
        ["Blast Furnace annual capacity", source.blast_furnace_annual, "tons/year"],
        ["Adjacent SKU setup", source.setup_adjacent_hours, "hours"],
        ["Non-adjacent SKU setup", source.setup_non_adjacent_hours, "hours"],
    ]
    add_rows(ws8, ["Input", "Value", "Unit"], rows, 3)
    start = 10
    ws8.cell(start, 1, "Rolling Mill Annual Capacities")
    ws8.cell(start + 1, 1, "Mill")
    for c, sku in enumerate(source.skus, start=2):
        ws8.cell(start + 1, c, sku)
    for r, mill in enumerate(["Rolling Mill 1", "Rolling Mill 2"], start=start + 2):
        ws8.cell(r, 1, mill)
        for c, sku in enumerate(source.skus, start=2):
            ws8.cell(r, c, source.mill_caps_annual[mill][sku])

    for sheet in wb.worksheets:
        input_fill, _ = style_sheet(sheet)
        sheet.freeze_panes = "A4"
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
        for cell in sheet[3]:
            if cell.value is not None:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
    ws["B10"].number_format = "0.0%"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def validate_output(path: Path) -> None:
    wb = load_workbook(path, data_only=True)
    required = [
        "Dashboard",
        "Demand_Used",
        "Optimized_MPS",
        "Campaign_Plan",
        "Capacity_Requirements",
        "MRP_Production_Req",
        "Raw_Material_MRP",
        "Given_Source_Data",
    ]
    missing = [name for name in required if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"Output workbook is missing sheets: {missing}")
    dash = wb["Dashboard"]
    if dash["B7"].value is None:
        raise ValueError("Dashboard verification failed: planned production is blank.")
    for ws in wb.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            raise ValueError(f"Output sheet appears blank: {ws.title}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Steel S&OP/MPS optimization agent")
    parser.add_argument("--template", default="Data Format (1).xlsx", help="Path to the instructor/template workbook")
    parser.add_argument("--input", default="Steel_SOP_MPS_Input.xlsx", help="Editable input workbook")
    parser.add_argument("--output", default="SOP_MPS_Optimized_Output.xlsx", help="Generated output workbook")
    parser.add_argument("--create-input-only", action="store_true", help="Create/update the editable input workbook and stop")
    args = parser.parse_args()

    template_path = Path(args.template)
    input_path = Path(args.input)
    output_path = Path(args.output)

    source = parse_template(template_path)
    if not input_path.exists():
        create_input_workbook(source, input_path)
        print(f"Created editable input workbook: {input_path}")

    if args.create_input_only:
        return

    demand, active, bom_factors = read_inputs(input_path, source)
    plans, campaigns, capacity_rows = run_optimization(source, demand, active)
    write_output(output_path, source, input_path, demand, active, bom_factors, plans, campaigns, capacity_rows)
    validate_output(output_path)

    print("Optimization completed.")
    print(f"Input workbook:  {input_path}")
    print(f"Output workbook: {output_path}")
    print(f"Total planned production: {sum(p.planned_production for p in plans):.2f} tons")
    print(f"Total unmet demand:       {sum(p.unmet_demand for p in plans):.2f} tons")
    print("Raw-material MRP was computed only if BOM factors were entered in BOM_Input_Optional.")


if __name__ == "__main__":
    main()
